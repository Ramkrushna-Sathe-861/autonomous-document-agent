"""Document generation tool using python-docx."""

import logging
import re
from collections.abc import Awaitable, Callable
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.config import get_settings
from app.core.exceptions import DocumentGenerationError
from app.schemas import ExecutionPlan, ExecutorResult

logger = logging.getLogger(__name__)


class DocumentTool:
    """Professional Word document generation and formatting."""

    def __init__(self):
        self.settings = get_settings()
        self.output_dir = self.settings.get_output_path()
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def build_node(self) -> Callable[[dict[str, object]], Awaitable[dict[str, object]]]:
        """Expose document rendering as a LangGraph node."""

        async def node(state: dict[str, object]) -> dict[str, object]:
            plan = state.get("plan")
            assert isinstance(plan, ExecutionPlan)
            sections = state.get("sections", [])
            assert isinstance(sections, list)
            filename = self._filename(plan.document_title)
            path = self.render_document(plan.document_title, plan.assumptions, sections, filename)
            return {"document_path": str(path)}

        return node

    def create_document(self, title: str, author: str = "", company: str = "") -> Document:
        try:
            doc = Document()
            core_props = doc.core_properties
            core_props.author = author or self.settings.document_author
            core_props.company = company or self.settings.document_company
            core_props.title = title

            title_paragraph = doc.add_paragraph(title)
            title_run = title_paragraph.runs[0]
            title_run.font.size = Pt(24)
            title_run.font.bold = True
            title_run.font.color.rgb = RGBColor(31, 78, 121)
            title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph()
            return doc
        except Exception as exc:
            logger.error("Failed to create document: %s", exc)
            raise DocumentGenerationError(f"Document creation failed: {exc}") from exc

    def add_heading(self, doc: Document, text: str, level: int = 1) -> None:
        try:
            heading = doc.add_heading(text, level=level)
            heading_run = heading.runs[0] if heading.runs else None
            if heading_run:
                if level == 1:
                    heading_run.font.color.rgb = RGBColor(31, 78, 121)
                    heading_run.font.size = Pt(18)
                elif level == 2:
                    heading_run.font.color.rgb = RGBColor(68, 114, 196)
                    heading_run.font.size = Pt(14)
                else:
                    heading_run.font.color.rgb = RGBColor(112, 144, 192)
                    heading_run.font.size = Pt(12)
        except Exception as exc:
            logger.error("Failed to add heading: %s", exc)
            raise DocumentGenerationError(f"Heading addition failed: {exc}") from exc

    def add_paragraph(self, doc: Document, text: str, style: Optional[str] = None) -> None:
        try:
            paragraph = doc.add_paragraph(text, style=style)
            for run in paragraph.runs:
                run.font.name = self.settings.document_font_name
                run.font.size = Pt(self.settings.document_font_size)
        except Exception as exc:
            logger.error("Failed to add paragraph: %s", exc)
            raise DocumentGenerationError(f"Paragraph addition failed: {exc}") from exc

    def add_bulleted_list(self, doc: Document, items: list[str]) -> None:
        try:
            for item in items:
                paragraph = doc.add_paragraph(item, style="List Bullet")
                for run in paragraph.runs:
                    run.font.name = self.settings.document_font_name
                    run.font.size = Pt(self.settings.document_font_size)
        except Exception as exc:
            logger.error("Failed to add bulleted list: %s", exc)
            raise DocumentGenerationError(f"Bulleted list addition failed: {exc}") from exc

    def add_numbered_list(self, doc: Document, items: list[str]) -> None:
        try:
            for item in items:
                paragraph = doc.add_paragraph(item, style="List Number")
                for run in paragraph.runs:
                    run.font.name = self.settings.document_font_name
                    run.font.size = Pt(self.settings.document_font_size)
        except Exception as exc:
            logger.error("Failed to add numbered list: %s", exc)
            raise DocumentGenerationError(f"Numbered list addition failed: {exc}") from exc

    def add_table(
        self, doc: Document, rows: int, cols: int, headers: Optional[list[str]] = None
    ) -> None:
        try:
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Table Grid"
            if headers:
                header_cells = table.rows[0].cells
                for index, header in enumerate(headers):
                    if index < len(header_cells):
                        header_cells[index].text = header
                        for paragraph in header_cells[index].paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True
        except Exception as exc:
            logger.error("Failed to add table: %s", exc)
            raise DocumentGenerationError(f"Table addition failed: {exc}") from exc

    def add_page_break(self, doc: Document) -> None:
        try:
            doc.add_page_break()
        except Exception as exc:
            logger.error("Failed to add page break: %s", exc)
            raise DocumentGenerationError(f"Page break addition failed: {exc}") from exc

    def render_document(self, title: str, assumptions: list[str], sections: list[ExecutorResult], filename: str) -> Path:
        document = self.create_document(title)
        if assumptions:
            self.add_heading(document, "Planning Assumptions", level=1)
            self.add_bulleted_list(document, assumptions)
        for section in sections:
            self.add_heading(document, section.section_title, level=1)
            self._add_content(document, section.content)
        return self.save_document(document, filename)

    def _add_content(self, document: Document, content: str) -> None:
        paragraphs: list[str] = []
        bullets: list[str] = []

        def flush() -> None:
            if paragraphs:
                self.add_paragraph(document, "\n".join(paragraphs))
                paragraphs.clear()
            if bullets:
                self.add_bulleted_list(document, bullets.copy())
                bullets.clear()

        for line in content.splitlines():
            stripped = line.strip()
            if stripped.startswith(("- ", "* ")):
                if paragraphs:
                    flush()
                bullets.append(stripped[2:].strip())
            elif not stripped:
                flush()
            else:
                if bullets:
                    flush()
                paragraphs.append(stripped)
        flush()

    @staticmethod
    def _filename(title: str) -> str:
        slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:50]
        return f"{slug or 'document'}-{__import__('uuid').uuid4().hex[:8]}.docx"

    def save_document(self, doc: Document, filename: str) -> Path:
        try:
            if not filename.endswith(".docx"):
                filename = f"{filename}.docx"
            file_path = self.output_dir / filename
            if file_path.exists():
                file_path.unlink()
            doc.save(str(file_path))
            logger.info("Document saved: %s", file_path)
            return file_path
        except Exception as exc:
            logger.error("Failed to save document: %s", exc)
            raise DocumentGenerationError(f"Document save failed: {exc}") from exc
